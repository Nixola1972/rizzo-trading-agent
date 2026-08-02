#!/usr/bin/env python3
"""Stampabili DOCX brandizzati Move 37.

    python build_docx.py --solo scheda-processo     # quella che si usa davvero
    python build_docx.py --solo canvas
    python build_docx.py --md out/esempio/verbale.md
    python build_docx.py --tutti

Due famiglie di output:

  · **moduli in bianco** — scheda processo A4 e canvas A3, generati da zero e
    pensati per essere riempiti a penna in riunione;
  · **documenti** — dossier, verbale e offerta, convertiti dal Markdown prodotto
    dagli altri tool.

Il meccanismo ricalca quello già collaudato per le offerte Trade Skill: un JSON o
un Markdown entra, un documento impaginato esce.
"""

import argparse
import pathlib
import re
import sys

try:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor
except ImportError:
    sys.exit("Serve python-docx: pip install -r requirements.txt")

BRAND = "MOVE 37"
GRIGIO = RGBColor(0x66, 0x66, 0x66)
NERO = RGBColor(0x1A, 0x1A, 0x1A)

CAMPI_SCHEDA = [
    ("① VOLUME", ["quante volte al mese"]),
    ("② CHI E QUANTO", ["chi ci mette le mani", "quanto tempo per volta"]),
    ("③ DOVE STANNO I DATI", ["dove", "formato / sistema", "chi li inserisce"]),
    ("④ QUANDO VA STORTO", ["cosa succede", "ogni quanto"]),
    ("⑤ CHI SI LAMENTA", ["☐ cliente   ☐ capo   ☐ chi esegue   ☐ nessuno ← attenzione"]),
    ("⑥ GIÀ PROVATO", ["cosa", "com'è andata"]),
    ("⑦ DECIDE / SI OPPONE", ["decide", "si oppone"]),
]

COLONNE_CANVAS = [
    "ARRIVA LA RICHIESTA",
    "FACCIAMO L'OFFERTA",
    "IL CLIENTE ACCETTA",
    "SI PRODUCE / ESEGUE",
    "SI CONSEGNA",
    "SI INCASSA",
]
RIGHE_CANVAS = ["chi", "quanto tempo", "quante volte/mese", "dove sono i dati", "cosa blocca ⚡"]


def _documento(orizzontale=False, margine=1.8):
    doc = Document()
    s = doc.sections[0]
    if orizzontale:
        s.orientation = WD_ORIENT.LANDSCAPE
        s.page_width, s.page_height = Cm(42.0), Cm(29.7)  # A3
        margine = 1.2
    else:
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)  # A4, non Letter
    for lato in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(s, lato, Cm(margine))
    normale = doc.styles["Normal"]
    normale.font.name = "Calibri"
    normale.font.size = Pt(10.5)
    normale.font.color.rgb = NERO
    return doc


def _intestazione(doc, titolo, sottotitolo=None):
    p = doc.add_paragraph()
    r = p.add_run(BRAND)
    r.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = GRIGIO
    p.paragraph_format.space_after = Pt(2)

    p = doc.add_paragraph()
    r = p.add_run(titolo)
    r.bold = True
    r.font.size = Pt(17)
    p.paragraph_format.space_after = Pt(2)

    if sottotitolo:
        p = doc.add_paragraph()
        r = p.add_run(sottotitolo)
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = GRIGIO
    return doc


def _riga_vuota(doc, etichetta="", larghezza=68):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    if etichetta:
        r = p.add_run(f"{etichetta}  ")
        r.font.size = Pt(9)
        r.font.color.rgb = GRIGIO
    r = p.add_run("." * larghezza)
    r.font.color.rgb = GRIGIO
    return p


def scheda_processo(percorso):
    doc = _documento()
    _intestazione(
        doc,
        "Scheda processo",
        "Una per processo. I campi si riempiono in ordine sparso, seguendo il cliente.",
    )
    doc.add_paragraph()

    _riga_vuota(doc, "PROCESSO (parole del cliente)", 55)
    _riga_vuota(doc, "Azienda", 30).add_run("     Data  " + "." * 18).font.color.rgb = GRIGIO
    doc.add_paragraph()

    for titolo, sotto in CAMPI_SCHEDA:
        p = doc.add_paragraph()
        r = p.add_run(titolo)
        r.bold = True
        r.font.size = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        for s in sotto:
            if s.startswith("☐"):
                q = doc.add_paragraph()
                q.paragraph_format.left_indent = Cm(0.6)
                q.paragraph_format.space_after = Pt(6)
                q.add_run(s).font.size = Pt(10)
            else:
                q = _riga_vuota(doc, s, 46)
                q.paragraph_format.left_indent = Cm(0.6)

    doc.add_paragraph()
    _riga_vuota(doc, "LEAD TIME (dall'inizio alla fine)", 30)
    _riga_vuota(doc, "TOUCH TIME (lavoro vero)", 34)

    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run("CITAZIONI TESTUALI")
    r.bold = True
    r.font.size = Pt(10)
    q = p.add_run("   (parole sue, virgolettate — vanno nel verbale così come sono)")
    q.font.size = Pt(8)
    q.font.color.rgb = GRIGIO
    for _ in range(3):
        _riga_vuota(doc, "", 76)

    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run("DA VEDERE / DA CHIEDERE DOPO")
    r.bold = True
    r.font.size = Pt(10)
    for _ in range(2):
        _riga_vuota(doc, "○", 74)

    doc.save(percorso)
    return percorso


def canvas(percorso):
    doc = _documento(orizzontale=True)
    _intestazione(
        doc,
        "Dal primo contatto al pagamento incassato",
        "Solo modalità B — se il cliente ha già un'agenda, questo foglio resta in cartella. "
        "Le colonne si rinominano seguendo lui.",
    )
    doc.add_paragraph()

    t = doc.add_table(rows=len(RIGHE_CANVAS) + 1, cols=len(COLONNE_CANVAS) + 1)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER

    t.cell(0, 0).text = ""
    for j, col in enumerate(COLONNE_CANVAS, 1):
        c = t.cell(0, j).paragraphs[0]
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.add_run(col)
        r.bold = True
        r.font.size = Pt(9)

    for i, riga in enumerate(RIGHE_CANVAS, 1):
        c = t.cell(i, 0).paragraphs[0]
        r = c.add_run(riga)
        r.bold = True
        r.font.size = Pt(9)
        for j in range(1, len(COLONNE_CANVAS) + 1):
            t.cell(i, j).paragraphs[0].add_run("\n\n")

    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run("LEAD TIME TOTALE  " + "." * 22 + "        di cui LAVORO VERO  " + "." * 22)
    r.bold = True
    r.font.size = Pt(11)

    p = doc.add_paragraph()
    r = p.add_run(
        "Si parte dalla fine («quando incassate?») e si risale. Il pennarello va in mano a loro. "
        "La riga «cosa blocca» è quella che conta: è lì che stanno i progetti."
    )
    r.italic = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = GRIGIO

    doc.save(percorso)
    return percorso


# --- conversione Markdown → DOCX -------------------------------------------------

_GRASSETTO = re.compile(r"\*\*(.+?)\*\*")


def _inline(paragrafo, testo, base=Pt(10.5)):
    """Rende il grassetto `**...**`; il resto passa come testo semplice."""
    for i, pezzo in enumerate(_GRASSETTO.split(testo)):
        if not pezzo:
            continue
        r = paragrafo.add_run(pezzo)
        r.bold = i % 2 == 1
        r.font.size = base


def _tabella(doc, righe):
    corpo = [r for r in righe if not re.fullmatch(r"\|[\s:|-]+\|", r.strip())]
    celle = [[c.strip() for c in r.strip().strip("|").split("|")] for r in corpo]
    if not celle:
        return
    larghezza = max(len(r) for r in celle)
    t = doc.add_table(rows=len(celle), cols=larghezza)
    t.style = "Table Grid"
    for i, riga in enumerate(celle):
        for j in range(larghezza):
            p = t.cell(i, j).paragraphs[0]
            _inline(p, riga[j] if j < len(riga) else "", Pt(9.5))
            if i == 0:
                for r in p.runs:
                    r.bold = True


def markdown(percorso_md, percorso_out):
    righe = pathlib.Path(percorso_md).read_text(encoding="utf-8").split("\n")
    doc = _documento()
    titolo_fatto = False
    i = 0
    while i < len(righe):
        riga = righe[i]
        vuota = riga.strip()

        if vuota.startswith("|"):
            blocco = []
            while i < len(righe) and righe[i].strip().startswith("|"):
                blocco.append(righe[i])
                i += 1
            _tabella(doc, blocco)
            doc.add_paragraph()
            continue

        if not vuota:
            i += 1
            continue

        if vuota.startswith("# ") and not titolo_fatto:
            _intestazione(doc, vuota[2:])
            titolo_fatto = True
        elif vuota.startswith("### "):
            doc.add_heading(vuota[4:], level=3)
        elif vuota.startswith("## "):
            doc.add_heading(vuota[3:], level=2)
        elif vuota.startswith("# "):
            doc.add_heading(vuota[2:], level=1)
        elif vuota.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            _inline(p, vuota[2:], Pt(10))
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = GRIGIO
        elif re.match(r"^[-*] ", vuota):
            p = doc.add_paragraph(style="List Bullet")
            _inline(p, vuota[2:])
        elif re.match(r"^\d+\. ", vuota):
            p = doc.add_paragraph(style="List Number")
            _inline(p, vuota.split(". ", 1)[1])
        elif set(vuota) <= {"-", "*", "_"} and len(vuota) >= 3:
            doc.add_paragraph()
        else:
            p = doc.add_paragraph()
            _inline(p, vuota.replace("*", ""))
        i += 1

    doc.save(percorso_out)
    return percorso_out


def main():
    ap = argparse.ArgumentParser(description="Stampabili DOCX Move 37")
    ap.add_argument(
        "--solo",
        choices=["scheda-processo", "canvas"],
        help="Genera un solo modulo in bianco",
    )
    ap.add_argument("--md", help="Converte un Markdown (dossier, verbale, offerta) in DOCX")
    ap.add_argument("--tutti", action="store_true", help="Genera entrambi i moduli in bianco")
    ap.add_argument("--out", default="stampabili")
    a = ap.parse_args()

    if not (a.solo or a.md or a.tutti):
        ap.error("indicare --solo, --md oppure --tutti")

    dest = pathlib.Path(a.out)
    dest.mkdir(parents=True, exist_ok=True)
    fatti = []

    if a.md:
        sorgente = pathlib.Path(a.md)
        fatti.append(markdown(sorgente, dest / f"{sorgente.stem}.docx"))
    if a.tutti or a.solo == "scheda-processo":
        fatti.append(scheda_processo(dest / "scheda-processo-A4.docx"))
    if a.tutti or a.solo == "canvas":
        fatti.append(canvas(dest / "canvas-mappa-processi-A3.docx"))

    for f in fatti:
        print(f)


if __name__ == "__main__":
    main()
